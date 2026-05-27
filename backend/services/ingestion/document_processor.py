"""
Document Processing Service
Extract text, chunk, and store documents
"""

import os
from typing import Dict
from sqlalchemy.orm import Session
import sys
from pathlib import Path

# Add backend directory to path
backend_dir = Path(__file__).parent.parent.parent
sys.path.insert(0, str(backend_dir))

from models.document import Document, Chunk
from config import settings


async def process_document(file_path: str, filename: str, db: Session) -> Dict:
    """
    Process uploaded document:
    1. Extract text
    2. Chunk content
    3. Generate embeddings
    4. Store in database and vector store
    
    Args:
        file_path: Path to uploaded file
        filename: Original filename
        db: Database session
    
    Returns:
        Dict with document info
    """
    # Get file info
    file_size = os.path.getsize(file_path)
    file_ext = os.path.splitext(filename)[1].lower()
    
    # Create document record
    document = Document(
        filename=filename,
        file_type=file_ext.replace(".", ""),
        file_size=file_size,
        file_path=file_path,
        status="processing"
    )
    db.add(document)
    db.commit()
    db.refresh(document)
    
    try:
        # 1. Extract text
        text = await extract_text(file_path, file_ext)
        
        # 2. Chunk text
        chunks = await chunk_text(text, filename)
        
        # 3. Store chunks
        chunk_count = 0
        for idx, chunk_data in enumerate(chunks):
            chunk = Chunk(
                document_id=document.id,
                content=chunk_data["content"],
                chunk_index=idx,
                page_number=chunk_data.get("page_number"),
                section_title=chunk_data.get("section_title"),
                extra_metadata=chunk_data.get("metadata", {})
            )
            db.add(chunk)
            chunk_count += 1
        
        # Update document status
        document.status = "completed"
        document.chunks_count = chunk_count
        db.commit()
        
        # 4. Generate embeddings and store in ChromaDB
        from services.retrieval.vector_store import add_to_vector_store
        await add_to_vector_store(chunks, document.id)
        
        return {
            "id": document.id,
            "filename": filename,
            "chunks_count": chunk_count,
            "status": "completed"
        }
    
    except Exception as e:
        document.status = "failed"
        db.commit()
        raise e


async def extract_text(file_path: str, file_ext: str) -> str:
    """
    Extract text from document
    """
    if file_ext == ".txt" or file_ext == ".md":
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    
    elif file_ext == ".pdf":
        import PyPDF2
        text = ""
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() + "\n"
        return text
    
    elif file_ext == ".docx":
        import docx
        doc = docx.Document(file_path)
        text = "\n".join([para.text for para in doc.paragraphs])
        return text
    
    elif file_ext == ".eml":
        from email import message_from_file
        with open(file_path, "r") as f:
            msg = message_from_file(f)
            text = ""
            if msg.is_multipart():
                for part in msg.walk():
                    if part.get_content_type() == "text/plain":
                        text += part.get_payload(decode=True).decode()
            else:
                text = msg.get_payload(decode=True).decode()
        return text
    
    else:
        raise ValueError(f"Unsupported file type: {file_ext}")


async def chunk_text(text: str, filename: str) -> list:
    """
    Chunk text into smaller pieces using RecursiveCharacterTextSplitter
    """
    from langchain.text_splitter import RecursiveCharacterTextSplitter
    
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        length_function=len,
        separators=["\n\n", "\n", " ", ""]
    )
    
    chunks_text = splitter.split_text(text)
    
    chunks = []
    for idx, chunk_content in enumerate(chunks_text):
        if chunk_content.strip():
            chunks.append({
                "content": chunk_content,
                "page_number": None,  # TODO: Track page numbers from PDF/DOCX
                "section_title": None,  # TODO: Extract section titles
                "metadata": {
                    "filename": filename,
                    "chunk_index": idx
                }
            })
    
    return chunks
